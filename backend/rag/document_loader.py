"""Document loader for PDF, DOCX, XLSX."""
import io
from pathlib import Path
from typing import Optional

import sys
sys.path.insert(0, str(Path(__file__).parent.parent.parent))


def load_document(file_path: str | Path) -> str:
    """Load document content as text. Supports PDF, DOCX, XLSX, TXT."""
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _load_pdf(path)
    elif suffix in (".docx", ".doc"):
        return _load_docx(path)
    elif suffix in (".xlsx", ".xls"):
        return _load_xlsx(path)
    elif suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type: {suffix}")


def _load_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    chunks = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            chunks.append(text)
    return "\n\n".join(chunks)


def _load_docx(path: Path) -> str:
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def _load_xlsx(path: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    chunks = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(c) if c is not None else "" for c in row)
            if row_text.strip():
                chunks.append(row_text)
    return "\n".join(chunks)


def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> list[str]:
    """Split text into overlapping chunks."""
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        if chunk.strip():
            chunks.append(chunk.strip())
        start = end - overlap
    return chunks
